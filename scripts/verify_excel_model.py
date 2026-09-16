"""Milestone 5: programmatic verification of the Excel executive model.

Opens the workbook with a real formula-evaluation engine (the `formulas`
package -- this environment's sandboxed LibreOffice cannot load ANY xlsx
file, confirmed with a trivial one-cell test file, so `formulas` is the
available substitute for "open, recalculate, inspect" here), recalculates
every formula, and checks:

1. No formula error values anywhere (#REF!, #VALUE!, #DIV/0!, etc.)
2. The live Scenario Forecast / Cash-Flow Bridge / Investment Capacity /
   Capital Allocation / DCF Valuation sheets reproduce target_cash.forecast
   and target_cash.valuation's own Python-computed numbers exactly, for
   whichever scenario the selector cell holds at save time (Base).
3. The Capital Allocation sheet's no-double-counting proof row reads "OK"
   for every year.

Usage: .venv/bin/python scripts/verify_excel_model.py
"""
import os
import sys

sys.path.insert(0, "src")

import formulas
import openpyxl

from target_cash import forecast as f
from target_cash import valuation as v

PATH = "deliverables/Target_Cash_Flow_Investment_Capacity_Model.xlsx"
SHEET_PREFIX = "'[Target_Cash_Flow_Investment_Capacity_Model.xlsx]"

errors_found = []


def check(condition, message):
    status = "PASS" if condition else "FAIL"
    print(f"[{status}] {message}")
    if not condition:
        errors_found.append(message)


print(f"Loading and recalculating {PATH} ...")
xl_model = formulas.ExcelModel().loads(PATH).finish()
solution = xl_model.calculate()
print(f"Recalculated {len(solution)} cells.\n")


def cell(sheet, ref):
    return solution[f"{SHEET_PREFIX}{sheet.upper()}'!{ref}"].value[0][0]


# --- 1. No formula errors anywhere ------------------------------------------
ERROR_TOKENS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!", "#N/A")
error_cells = []
for key, res in solution.items():
    try:
        val = res.value[0][0]
    except Exception:
        continue
    if isinstance(val, str) and any(tok in val for tok in ERROR_TOKENS):
        error_cells.append((key, val))
check(len(error_cells) == 0, f"No formula errors anywhere in the workbook (found {len(error_cells)})")
for k, v_ in error_cells[:20]:
    print(f"   ERROR CELL: {k} = {v_}")

# --- 2. Live sheets reproduce Python exactly (Base scenario, as saved) ------
assumptions = f.build_assumptions()
forecasts = f.run_all_scenarios(assumptions)
years = forecasts["base"]
FY_COLS = ["D", "E", "F", "G", "H"]

wb_ro = openpyxl.load_workbook(PATH)
ws_fc = wb_ro["Scenario Forecast"]
ws_cf = wb_ro["Cash-Flow Bridge"]
ws_ic = wb_ro["Investment Capacity"]
ws_ca = wb_ro["Capital Allocation"]
ws_dcf = wb_ro["DCF Valuation"]


def row_of(ws, label):
    for row in ws.iter_rows(min_col=1, max_col=1):
        for c in row:
            if c.value == label:
                return c.row
    raise KeyError(label)


rev_row = row_of(ws_fc, "Revenue")
ni_row = row_of(ws_fc, "Net Income")
eps_row = row_of(ws_fc, "Diluted EPS ($)")
cfo_row = row_of(ws_cf, "Cash Flow from Operations (CFO)")
fcf_row = row_of(ws_cf, "Free Cash Flow (CFO - CapEx)")
end_cash_row = row_of(ws_cf, "Ending Cash")
dep_row = row_of(ws_ic, "DEPLOYABLE CAPACITY")
wf_end_row = row_of(ws_ca, "8. = Ending Cash")
proof_row = row_of(ws_ca, "Ending Cash (Sheet 8) matches Step 8 above?")

all_match = True
for i, (col, y) in enumerate(zip(FY_COLS, years)):
    excel_rev = cell("SCENARIO FORECAST", f"{col}{rev_row}")
    excel_ni = cell("SCENARIO FORECAST", f"{col}{ni_row}")
    excel_eps = cell("SCENARIO FORECAST", f"{col}{eps_row}")
    excel_cfo = cell("CASH-FLOW BRIDGE", f"{col}{cfo_row}")
    excel_fcf = cell("CASH-FLOW BRIDGE", f"{col}{fcf_row}")
    excel_end_cash = cell("CASH-FLOW BRIDGE", f"{col}{end_cash_row}")
    excel_dep = cell("INVESTMENT CAPACITY", f"{col}{dep_row}")
    excel_wf_end = cell("CAPITAL ALLOCATION", f"{col}{wf_end_row}")
    proof_val = cell("CAPITAL ALLOCATION", f"{col}{proof_row}")

    matches = (
        abs(excel_rev - y.revenue) < 0.01 and abs(excel_ni - y.net_income) < 0.01
        and abs(excel_eps - y.diluted_eps) < 0.001 and abs(excel_cfo - y.operating_cash_flow) < 0.01
        and abs(excel_fcf - y.free_cash_flow) < 0.01 and abs(excel_end_cash - y.ending_cash) < 0.01
        and abs(excel_dep - y.deployable_capacity) < 0.01 and abs(excel_wf_end - y.ending_cash) < 0.01
        and proof_val == "OK"
    )
    all_match &= matches
    check(matches, f"FY{y.fiscal_year} Base scenario: Excel matches Python exactly "
                    f"(revenue, net income, EPS, CFO, FCF, ending cash, deployable capacity, waterfall proof)")

# --- 3. DCF sheet matches Python -------------------------------------------
val_assumptions = v.build_valuation_assumptions()
dcf_result = v.run_dcf(years, val_assumptions)
wacc_row = row_of(ws_dcf, "WACC")
ev_row = row_of(ws_dcf, "Enterprise Value")
equity_row = row_of(ws_dcf, "Equity Value")
per_share_row = row_of(ws_dcf, "IMPLIED VALUE PER SHARE")

excel_wacc = cell("DCF VALUATION", f"B{wacc_row}")
excel_ev = cell("DCF VALUATION", f"B{ev_row}")
excel_equity = cell("DCF VALUATION", f"B{equity_row}")
excel_per_share = cell("DCF VALUATION", f"B{per_share_row}")

check(abs(excel_wacc * 100 - dcf_result.wacc_pct) < 0.001, "DCF: WACC matches Python")
check(abs(excel_ev - dcf_result.enterprise_value) < 0.1, "DCF: Enterprise Value matches Python")
check(abs(excel_equity - dcf_result.equity_value) < 0.1, "DCF: Equity Value matches Python")
check(abs(excel_per_share - dcf_result.implied_value_per_share) < 0.01, "DCF: Implied Value/Share matches Python")

# --- 4. Structural checks ---------------------------------------------------
EXPECTED_SHEETS = [
    "Cover & Instructions", "Executive Summary", "Historical Financials", "Filing-Vintage Comparison",
    "Quarterly Cash Proof", "Forecast Assumptions", "Scenario Forecast", "Cash-Flow Bridge",
    "Investment Capacity", "Capital Allocation", "DCF Valuation", "Sensitivities", "Source & Lineage",
    "Validation Summary", "Limitations",
]
check(wb_ro.sheetnames == EXPECTED_SHEETS, f"All 15 sheets present in the required order (got {wb_ro.sheetnames})")

ws_asm = wb_ro["Forecast Assumptions"]
dv_found = any(dv.type == "list" for dv in ws_asm.data_validations.dataValidation)
check(dv_found, "Scenario selector dropdown (data validation) present on Forecast Assumptions!C2")

for name in ["Scenario Forecast", "Cash-Flow Bridge", "Investment Capacity", "Forecast Assumptions"]:
    ws = wb_ro[name]
    check(ws.freeze_panes is not None, f"{name}: frozen panes set")

# --- 5. Scenario selector genuinely recalculates (Upside/Downside) --------
import tempfile

for scenario_label, scenario_key in [("Upside", "upside"), ("Downside", "downside")]:
    wb_switch = openpyxl.load_workbook(PATH)
    wb_switch["Forecast Assumptions"]["C2"] = scenario_label
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name
    wb_switch.save(tmp_path)
    tmp_name = tmp_path.split("/")[-1]
    xl_model_2 = formulas.ExcelModel().loads(tmp_path).finish()
    solution_2 = xl_model_2.calculate()
    excel_rev_2030 = solution_2[f"'[{tmp_name}]SCENARIO FORECAST'!H{rev_row}"].value[0][0]
    excel_eps_2030 = solution_2[f"'[{tmp_name}]SCENARIO FORECAST'!H{eps_row}"].value[0][0]
    py_years = forecasts[scenario_key]
    check(
        abs(excel_rev_2030 - py_years[-1].revenue) < 0.01 and abs(excel_eps_2030 - py_years[-1].diluted_eps) < 0.001,
        f"Scenario selector -> {scenario_label}: FY2030 revenue and EPS recalculate to match Python exactly",
    )
    os.remove(tmp_path)

print()
if errors_found:
    print(f"VERIFICATION FAILED: {len(errors_found)} issue(s) found.")
    sys.exit(1)
else:
    print("ALL VERIFICATION CHECKS PASSED.")
